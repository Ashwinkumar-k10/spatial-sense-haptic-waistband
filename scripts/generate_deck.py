"""
SPATIAL-SENSE: MASTER POWERPOINT GENERATION SCRIPT
Theme 2: Assistive Solutions (EnAble India), IIT Indore
Vishwakarma Awards 2026–27 | Team TECHTONICS

Generates the complete 14-slide professional presentation matching Theme 7 design language
using Theme 2 technical content.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.500000)
    blank_layout = prs.slide_layouts[6]

    # Institutional Color Palette
    NAVY_PRIMARY = RGBColor(26, 54, 112)       # #1A3670 - Deep Navy
    NAVY_DARK = RGBColor(15, 23, 42)           # #0F172A - Dark Slate
    BLUE_ACCENT = RGBColor(0, 112, 174)        # #0070AE - Cerulean Blue
    BLUE_CARD_BG = RGBColor(241, 245, 249)     # #F1F5F9 - Pale Slate Blue
    BLUE_CARD_BORDER = RGBColor(203, 213, 225) # #CBD5E1 - Border Gray
    GREEN_ACCENT = RGBColor(22, 101, 52)       # #166534 - Solution Green
    GREEN_CARD_BG = RGBColor(240, 253, 244)    # #F0FDF4 - Pale Green
    GREEN_CARD_BORDER = RGBColor(187, 247, 208)# #BBF7D0 - Border Green
    AMBER_ACCENT = RGBColor(217, 119, 6)       # #D97706 - Safety Amber
    RED_ACCENT = RGBColor(220, 38, 38)         # #DC2626 - Alert Red
    WHITE = RGBColor(255, 255, 255)            # #FFFFFF
    TEXT_DARK = RGBColor(30, 41, 59)           # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)       # #64748B

    FONT_TITLE = 'Georgia'
    FONT_BODY = 'Aptos'

    def add_header(slide, title_text, badge_text=""):
        if os.path.exists("assets/logo_maker_bhavan.png"):
            slide.shapes.add_picture("assets/logo_maker_bhavan.png", Inches(0.4), Inches(0.18), height=Inches(0.42))
        if os.path.exists("assets/logo_vishwakarma.png"):
            slide.shapes.add_picture("assets/logo_vishwakarma.png", Inches(1.35), Inches(0.12), height=Inches(0.55))
            
        tx_box = slide.shapes.add_textbox(Inches(2.25), Inches(0.15), Inches(7.3), Inches(0.55))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text.upper()
        p.font.name = FONT_TITLE
        p.font.size = Pt(19)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        
        if badge_text:
            badge_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, 
                Inches(9.6), Inches(0.18), Inches(3.3), Inches(0.42)
            )
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
            badge_box.line.color.rgb = RGBColor(191, 219, 254)
            badge_box.line.width = Pt(1)
            btf = badge_box.text_frame
            btf.vertical_anchor = MSO_ANCHOR.MIDDLE
            btf.word_wrap = True
            bp = btf.paragraphs[0]
            bp.alignment = PP_ALIGN.CENTER
            bp.text = badge_text.upper()
            bp.font.name = FONT_BODY
            bp.font.size = Pt(9.0)
            bp.font.bold = True
            bp.font.color.rgb = BLUE_ACCENT
            
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(0.4), Inches(0.72), Inches(12.533), Inches(0.015)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(226, 232, 240)
        line.line.fill.background()

    def create_card(slide, left, top, width, height, bg_color=WHITE, border_color=BLUE_CARD_BORDER, line_width=1):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(line_width)
        return shape

    def style_table(table, header_bg=NAVY_PRIMARY, alt_bg=RGBColor(248, 250, 252)):
        for col_idx in range(len(table.columns)):
            cell = table.cell(0, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_bg
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(8.0)
                p.font.bold = True
                p.font.color.rgb = WHITE
                
        for row_idx in range(1, len(table.rows)):
            row_bg = alt_bg if row_idx % 2 == 1 else WHITE
            for col_idx in range(len(table.columns)):
                cell = table.cell(row_idx, col_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_bg
                cell.margin_left = cell.margin_right = Inches(0.07)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                for p in cell.text_frame.paragraphs:
                    p.font.name = FONT_BODY
                    p.font.size = Pt(7.5)
                    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    if os.path.exists("assets/logo_maker_bhavan.png"):
        slide1.shapes.add_picture("assets/logo_maker_bhavan.png", Inches(0.6), Inches(0.38), height=Inches(0.68))
    if os.path.exists("assets/logo_vishwakarma.png"):
        slide1.shapes.add_picture("assets/logo_vishwakarma.png", Inches(2.2), Inches(0.3), height=Inches(0.85))

    banner_box = slide1.shapes.add_textbox(Inches(3.8), Inches(0.38), Inches(8.8), Inches(0.7))
    btf = banner_box.text_frame
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    bp.text = "VISHWAKARMA AWARDS 2026–27"
    bp.font.name = FONT_TITLE
    bp.font.size = Pt(22)
    bp.font.bold = True
    bp.font.color.rgb = NAVY_PRIMARY

    bp2 = btf.add_paragraph()
    bp2.text = "STAGE 2 SUBMISSION • NATIONAL INNOVATION & DESIGN COMPETITION"
    bp2.font.name = FONT_BODY
    bp2.font.size = Pt(10)
    bp2.font.bold = True
    bp2.font.color.rgb = BLUE_ACCENT

    left_card = create_card(slide1, Inches(0.6), Inches(1.35), Inches(6.0), Inches(5.65), bg_color=RGBColor(248, 250, 252), border_color=BLUE_CARD_BORDER, line_width=1.5)
    ltf = left_card.text_frame
    ltf.margin_left = ltf.margin_right = Inches(0.3)
    ltf.margin_top = Inches(0.25)
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "TRACK & INSTITUTION"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p = ltf.add_paragraph()
    p.text = "THEME 2 — ASSISTIVE SOLUTIONS (ENABLE INDIA)\nIIT INDORE"
    p.font.name = FONT_TITLE
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = ltf.add_paragraph()
    p.text = "\nTEAM IDENTITY"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p = ltf.add_paragraph()
    p.text = "TEAM TECHTONICS"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    p = ltf.add_paragraph()
    p.text = "\nPROJECT TITLE"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p = ltf.add_paragraph()
    p.text = "SPATIAL-SENSE"
    p.font.name = FONT_TITLE
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = ltf.add_paragraph()
    p.text = "A 360° Wearable Haptic-Sonar Waistband for Upper-Torso and Head-Level Collision Protection in Deaf-Blind Mobility"
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = ltf.add_paragraph()
    p.text = "\nBENEFICIARY FOCUS & ACCESSIBILITY"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p = ltf.add_paragraph()
    p.text = "• Primary: Deaf-blind individuals (over 450,000 in India)\n• Secondary: Visually impaired persons with age-related hearing loss\n• Tertiary: Orientation & Mobility (O&M) trainers (EnAble India)"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_DARK

    if os.path.exists("assets/spatial_sense_cover.jpg"):
        slide1.shapes.add_picture("assets/spatial_sense_cover.jpg", Inches(6.8), Inches(1.35), width=Inches(5.9), height=Inches(3.45))
    if os.path.exists("assets/dual_plane_geometry.png"):
        slide1.shapes.add_picture("assets/dual_plane_geometry.png", Inches(6.8), Inches(4.9), width=Inches(5.9), height=Inches(1.65))

    cap_card = create_card(slide1, Inches(6.8), Inches(6.6), Inches(5.9), Inches(0.4), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
    ctf = cap_card.text_frame
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    cp = ctf.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    cp.text = "HANDS-FREE 360° TACTILE OBSTACLE AWARENESS • DUAL-ELEVATION SENSING"
    cp.font.name = FONT_BODY
    cp.font.size = Pt(9.0)
    cp.font.bold = True
    cp.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "PROBLEM STATEMENT", "NEED FOR SPATIAL-SENSE")

    ov_card = create_card(slide2, Inches(0.4), Inches(0.85), Inches(7.5), Inches(1.85), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    otf = ov_card.text_frame
    otf.margin_left = otf.margin_right = otf.margin_top = Inches(0.18)
    otf.word_wrap = True
    p = otf.paragraphs[0]
    p.text = "PROJECT OVERVIEW & CORE PAIN POINT"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p = otf.add_paragraph()
    p.text = "Project Name: SPATIAL-SENSE | Track: Assistive Solutions (EnAble India), IIT Indore"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = otf.add_paragraph()
    p.text = 'Core Problem Statement: "Individuals with deaf-blindness lack non-visual, non-auditory spatial navigation tools, leaving them unprotected against upper-torso and head-level obstacles—such as low-hanging tree branches, truck cargo beds, open casement windows, and street awnings—which ground-sweeping white canes cannot detect."'
    p.font.name = FONT_BODY
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_DARK

    p = otf.add_paragraph()
    p.text = "Clinical Context: Over 450,000 deaf-blind individuals in India experience severe sensorineural hearing loss alongside visual impairment. Standard white canes sweep the ground plane, leaving the cranium and upper torso vulnerable. Acoustic smart canes fail because audio cues cannot be perceived. Frequent collisions cause blunt-force trauma, disorientation, and loss of walking independence."
    p.font.name = FONT_BODY
    p.font.size = Pt(8.0)
    p.font.color.rgb = TEXT_MUTED

    cmp_banner = create_card(slide2, Inches(0.4), Inches(2.8), Inches(7.5), Inches(0.4), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
    ctf = cmp_banner.text_frame
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = ctf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "THE EXISTING MOBILITY PROBLEM vs THE SPATIAL-SENSE SOLUTION"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = WHITE

    pills = [
        ("AI OPTICAL ToF", "Dual-plane distance sensing\n(0° waist & +25° head level)"),
        ("360° RADAR", "Full torso perimeter coverage\nvia 8 radial sensor nodes"),
        ("RESONANT HAPTICS", "205 Hz tactile distance-to-cadence\nintuitive abdominal pulses"),
        ("HANDS-FREE BELT", "Ergonomic wearable waistband;\npreserves cane for ground path")
    ]
    for idx, (p_title, p_desc) in enumerate(pills):
        px = Inches(0.4 + idx * 1.9)
        pc = create_card(slide2, px, Inches(3.28), Inches(1.8), Inches(1.15), bg_color=GREEN_CARD_BG, border_color=GREEN_CARD_BORDER)
        ptf = pc.text_frame
        ptf.margin_left = ptf.margin_right = ptf.margin_top = Inches(0.08)
        ptf.word_wrap = True
        p = ptf.paragraphs[0]
        p.text = p_title
        p.font.name = FONT_BODY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = GREEN_ACCENT
        p = ptf.add_paragraph()
        p.text = p_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    if os.path.exists("assets/theme2_before_cane.jpg"):
        slide2.shapes.add_picture("assets/theme2_before_cane.jpg", Inches(0.4), Inches(4.55), width=Inches(7.5), height=Inches(1.4))

    right_cards = [
        ("✦ Undetected Upper-Body Hazards", 
         "Standard white canes only sweep the ground plane, missing overhead tree branches, commercial shop awnings, open casement windows, and projecting truck cargo beds.", AMBER_ACCENT),
        ("● Audio Inaccessibility", 
         "Acoustic smart canes rely on audio beeps or bone conduction, which completely fail for deaf-blind individuals suffering profound sensorineural hearing loss.", RED_ACCENT),
        ("◎ Limited Spatial Awareness", 
         "Existing handheld assistive devices only cover a narrow forward cone (1.0–1.5 m); they do not provide 360° peripheral awareness around the user's torso.", BLUE_ACCENT),
        ("! Reduced Mobility Confidence", 
         "Frequent upper-body collisions cause facial lacerations, blunt-force head trauma, severe hesitation, and complete withdrawal from outdoor independent mobility.", RED_ACCENT),
        ("₹ Accessibility & High Cost Barrier", 
         "Imported haptic vests cost ₹1,20,000+ and are bulky; an affordable (₹9,500 prototype / ₹4,500 retail target), low-power wearable belt is urgently needed.", GREEN_ACCENT)
    ]

    for idx, (rc_title, rc_body, rc_color) in enumerate(right_cards):
        ry = Inches(0.85 + idx * 1.02)
        rc = create_card(slide2, Inches(8.1), ry, Inches(4.8), Inches(0.95), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
        rtf = rc.text_frame
        rtf.margin_left = rtf.margin_right = rtf.margin_top = Inches(0.12)
        rtf.word_wrap = True
        p = rtf.paragraphs[0]
        p.text = rc_title
        p.font.name = FONT_BODY
        p.font.size = Pt(8.8)
        p.font.bold = True
        p.font.color.rgb = rc_color
        p = rtf.add_paragraph()
        p.text = rc_body
        p.font.name = FONT_BODY
        p.font.size = Pt(7.8)
        p.font.color.rgb = TEXT_DARK

    b_cards = [
        ("DEAF-BLIND INDIVIDUALS", "Over 450,000 in India needing non-visual, non-auditory collision protection."),
        ("VISUALLY & HEARING IMPAIRED", "Adults with combined low-vision and age-related presbycusis."),
        ("O&M MOBILITY TRAINERS", "EnAble India & Blind Associations training independent navigation.")
    ]
    for idx, (b_title, b_desc) in enumerate(b_cards):
        bx = Inches(0.4 + idx * 2.5)
        bc = create_card(slide2, bx, Inches(6.05), Inches(2.4), Inches(1.25), bg_color=RGBColor(241, 245, 249), border_color=BLUE_CARD_BORDER)
        btf = bc.text_frame
        btf.margin_left = btf.margin_right = btf.margin_top = Inches(0.1)
        btf.word_wrap = True
        p = btf.paragraphs[0]
        p.text = b_title
        p.font.name = FONT_BODY
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = btf.add_paragraph()
        p.text = b_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    unc = create_card(slide2, Inches(7.9), Inches(6.05), Inches(5.0), Inches(1.25), bg_color=RGBColor(239, 246, 255), border_color=BLUE_ACCENT)
    utf = unc.text_frame
    utf.margin_left = utf.margin_right = utf.margin_top = Inches(0.12)
    utf.word_wrap = True
    p = utf.paragraphs[0]
    p.text = "THE UNMET NEED STATEMENT"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p = utf.add_paragraph()
    p.text = "A wearable, hands-free assistive device that extends obstacle detection beyond the ground plane to protect the upper torso and head, delivers non-auditory directional haptics, and supplements the white cane without sensory overload."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.5)
    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 3: PROPOSED SOLUTION
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "PROPOSED SOLUTION", "INNOVATION & NOVELTY")

    sub_box = slide3.shapes.add_textbox(Inches(0.4), Inches(0.78), Inches(12.5), Inches(0.35))
    stf = sub_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_top = 0
    sp = stf.paragraphs[0]
    sp.text = "SPATIAL-SENSE: 360° Wearable Haptic-Sonar Assistive System for Upper-Torso & Head-Level Collision Protection"
    sp.font.name = FONT_TITLE
    sp.font.size = Pt(11)
    sp.font.bold = True
    sp.font.color.rgb = NAVY_PRIMARY

    if os.path.exists("assets/spatial_sense_worn.jpg"):
        slide3.shapes.add_picture("assets/spatial_sense_worn.jpg", Inches(2.3), Inches(1.2), width=Inches(3.4), height=Inches(4.6))

    left_callouts = [
        ("8x VL53L1X ToF SENSORS", "4x at 0° waist level + 4x at +25° head level; 940nm Class-1 eye-safe LiDAR."),
        ("ESP32-S3 CONTROLLER", "Dual-core 240MHz MCU; Core 0 sensor DSP, Core 1 haptics, ULP sleep manager."),
        ("TCA9548A I²C MUX", "8-channel multiplexer; executes interleaved 4ms phase firing schedule."),
        ("BMI160 6-DOF IMU", "6-axis motion sensing; enables 15µA light sleep and 5ms step wake-up.")
    ]
    for idx, (c_title, c_desc) in enumerate(left_callouts):
        cy = Inches(1.2 + idx * 1.15)
        cc = create_card(slide3, Inches(0.4), cy, Inches(1.85), Inches(1.05), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
        ctf = cc.text_frame
        ctf.margin_left = ctf.margin_right = ctf.margin_top = Inches(0.08)
        ctf.word_wrap = True
        p = ctf.paragraphs[0]
        p.text = c_title
        p.font.name = FONT_BODY
        p.font.size = Pt(7.8)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = ctf.add_paragraph()
        p.text = c_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(6.8)
        p.font.color.rgb = TEXT_DARK

    right_callouts = [
        ("8x DRV2605L DRIVERS", "Dedicated I²C bus control, auto-resonance calibration, 150mA current limit."),
        ("8x 205Hz LRAs", "Linear resonant coin actuators; 2.2G peak tactile acceleration; 5ms rise time."),
        ("MODULAR TPU PODS", "Silicone-lined IP54 pods; ergonomic neoprene belt; skin temp <34.5°C."),
        ("3.7V 2500mAh Li-ion", "Hardware BMS with OVP/UVP/OCP; accessible USB-C charging interface.")
    ]
    for idx, (c_title, c_desc) in enumerate(right_callouts):
        cy = Inches(1.2 + idx * 1.15)
        cc = create_card(slide3, Inches(5.75), cy, Inches(1.85), Inches(1.05), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
        ctf = cc.text_frame
        ctf.margin_left = ctf.margin_right = ctf.margin_top = Inches(0.08)
        ctf.word_wrap = True
        p = ctf.paragraphs[0]
        p.text = c_title
        p.font.name = FONT_BODY
        p.font.size = Pt(7.8)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = ctf.add_paragraph()
        p.text = c_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(6.8)
        p.font.color.rgb = TEXT_DARK

    innovations = [
        ("1. Dual-Elevation Spatial Sensing", 
         "Four horizontal sensors at 0° scan parallel to ground for waist-level obstacles (walls, pedestrians), while four alternating sensors tilted upward at +25° project a protective cone up to 2.1 m height to detect overhead tree limbs, awnings, and signs.", BLUE_ACCENT),
        ("2. Full 360° Surrounding Tactile Perimeter", 
         "Eight sensors spaced at 45° intervals provide complete surround obstacle awareness, mapping environmental vectors directly to corresponding localized haptic nodes on the user's torso.", GREEN_ACCENT),
        ("3. Deterministic Hard Real-Time Processing", 
         "The ESP32-S3 performs non-blocking I²C sensor scheduling, median distance filtering, and haptic modulation in deterministic DSP—achieving <15 ms end-to-end latency without latency-heavy neural networks.", NAVY_PRIMARY),
        ("4. Resonant Haptic Actuation & Distance Cadence", 
         "Eight DRV2605L-driven LRAs operate strictly at mechanical resonance (205 Hz ± 3 Hz) for crisp tactile transients. Distance maps to pulse repetition cadence: 1.2 s pulse at 2.5 m accelerating smoothly to a 12 Hz urgent burst at 0.4 m.", AMBER_ACCENT),
        ("5. Accessibility-First Offline Operation", 
         "100% non-visual, non-auditory primary alerts. Features a zero-sight hardware self-test (double-tap triggers clockwise sweep) and an offline BLE companion app with TalkBack/VoiceOver and O&M trainer mirror.", BLUE_ACCENT)
    ]

    for idx, (i_title, i_desc, i_col) in enumerate(innovations):
        iy = Inches(1.2 + idx * 0.94)
        ic = create_card(slide3, Inches(7.8), iy, Inches(5.1), Inches(0.88), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
        itf = ic.text_frame
        itf.margin_left = itf.margin_right = itf.margin_top = Inches(0.12)
        itf.word_wrap = True
        p = itf.paragraphs[0]
        p.text = i_title
        p.font.name = FONT_BODY
        p.font.size = Pt(8.8)
        p.font.bold = True
        p.font.color.rgb = i_col
        p = itf.add_paragraph()
        p.text = i_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.5)
        p.font.color.rgb = TEXT_DARK

    bot_quote = create_card(slide3, Inches(0.4), Inches(6.0), Inches(12.5), Inches(0.85), bg_color=RGBColor(241, 245, 249), border_color=BLUE_ACCENT)
    qtf = bot_quote.text_frame
    qtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    qtf.margin_left = qtf.margin_right = Inches(0.2)
    qp = qtf.paragraphs[0]
    qp.text = '“SPATIAL-SENSE extends obstacle awareness beyond the ground plane by combining dual-elevation optical sensing with spatially localized tactile feedback—enabling hands-free, dignified independent navigation for deaf-blind individuals.”'
    qp.font.name = FONT_TITLE
    qp.font.size = Pt(10)
    qp.font.bold = True
    qp.font.color.rgb = NAVY_PRIMARY
    qp = qtf.add_paragraph()
    qp.text = "Note: Wearable illustrations represent proposed concept renders; functional specifications and performance metrics reflect engineering targets to be validated."
    qp.font.name = FONT_BODY
    qp.font.size = Pt(7.2)
    qp.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 4: USER JOURNEY
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "USER JOURNEY", "BEFORE-TO-AFTER MOBILITY IMPACT")

    sub_box = slide4.shapes.add_textbox(Inches(0.4), Inches(0.75), Inches(12.5), Inches(0.3))
    stf = sub_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_top = 0
    sp = stf.paragraphs[0]
    sp.text = "FROM GROUND-LEVEL CANE DETECTION TO 360° TACTILE OBSTACLE AWARENESS"
    sp.font.name = FONT_BODY
    sp.font.size = Pt(9.5)
    sp.font.bold = True
    sp.font.color.rgb = BLUE_ACCENT

    metrics = [
        ("8 ToF SENSORS", "Dual elevation planes\n(0° waist + 25° head level)"),
        ("360° PERIMETER", "Complete surround coverage\nvia 8 radial body nodes"),
        ("15 Hz REFRESH", "Interleaved 4ms phase offsets\nzero optical crosstalk"),
        ("< 15 ms LATENCY", "Deterministic sensor-to-haptic\nreal-time DSP response")
    ]
    for idx, (m_val, m_lbl) in enumerate(metrics):
        mx = Inches(0.4 + idx * 2.25)
        mc = create_card(slide4, mx, Inches(1.1), Inches(2.1), Inches(0.85), bg_color=RGBColor(248, 250, 252), border_color=BLUE_CARD_BORDER)
        mtf = mc.text_frame
        mtf.margin_left = mtf.margin_right = mtf.margin_top = Inches(0.08)
        mtf.word_wrap = True
        p = mtf.paragraphs[0]
        p.text = m_val
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = mtf.add_paragraph()
        p.text = m_lbl
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    b_card = create_card(slide4, Inches(0.4), Inches(2.05), Inches(4.3), Inches(3.6), bg_color=WHITE, border_color=RED_ACCENT)
    btf = b_card.text_frame
    btf.margin_left = btf.margin_right = btf.margin_top = Inches(0.12)
    btf.word_wrap = True
    p = btf.paragraphs[0]
    p.text = "BEFORE | CONVENTIONAL WHITE CANE ALONE"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT

    if os.path.exists("assets/theme2_before_cane.jpg"):
        slide4.shapes.add_picture("assets/theme2_before_cane.jpg", Inches(0.55), Inches(2.35), width=Inches(4.0), height=Inches(1.8))

    p = btf.add_paragraph()
    p.text = "\n\n\n\n\n\n\n• Scenario: Navigating a busy footpath. 15m ahead, an iron shop awning hangs at 1.4m chest height, with a low tree limb at 1.75m.\n• Limitation: Cane tip sweeps pavement (clear). Upper torso & forehead completely unprotected.\n• Result: Forehead strikes rigid metal corner at 0m. Facial laceration, blunt trauma, severe disorientation, and loss of walking confidence."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.5)
    p.font.color.rgb = TEXT_DARK

    a_card = create_card(slide4, Inches(4.85), Inches(2.05), Inches(4.3), Inches(3.6), bg_color=WHITE, border_color=GREEN_ACCENT)
    atf = a_card.text_frame
    atf.margin_left = atf.margin_right = atf.margin_top = Inches(0.12)
    atf.word_wrap = True
    p = atf.paragraphs[0]
    p.text = "AFTER | SPATIAL-SENSE + WHITE CANE (COMPLEMENTARY)"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    if os.path.exists("assets/theme2_after_spatial_sense.jpg"):
        slide4.shapes.add_picture("assets/theme2_after_spatial_sense.jpg", Inches(5.0), Inches(2.35), width=Inches(4.0), height=Inches(1.8))

    p = atf.add_paragraph()
    p.text = "\n\n\n\n\n\n\n• Early Detection: At 2.5m, the Front-Right +25° ToF sensor detects the iron awning.\n• Tactile Alert: Front-Right LRA delivers gentle 1 Hz pulses to the abdomen.\n• Safe Evasion: User extends a protective hand. At 0.8m, cadence accelerates to urgent 8-12 Hz burst. User ducks smoothly around obstacle without bodily impact."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.5)
    p.font.color.rgb = TEXT_DARK

    r_panel = create_card(slide4, Inches(9.3), Inches(1.1), Inches(3.6), Inches(4.55), bg_color=RGBColor(248, 250, 252), border_color=BLUE_CARD_BORDER)
    rtf = r_panel.text_frame
    rtf.margin_left = rtf.margin_right = rtf.margin_top = Inches(0.15)
    rtf.word_wrap = True

    p = rtf.paragraphs[0]
    p.text = "TARGET USER PROFILE"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p = rtf.add_paragraph()
    p.text = "Deaf-blind individuals, persons with low vision & hearing impairment, and O&M mobility trainers."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_DARK

    p = rtf.add_paragraph()
    p.text = "\nCURRENT MOBILITY LIMITATION"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT

    p = rtf.add_paragraph()
    p.text = "Cane detects only ground obstacles (<0.5m). Upper body strikes average 3.4 collisions per 100m in unfamiliar corridors at baseline walking speed of 0.52 m/s."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_DARK

    p = rtf.add_paragraph()
    p.text = "\nPROPOSED ACCESSIBILITY ADVANCE"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    p = rtf.add_paragraph()
    p.text = "Hands-free, zero-audio tactile radar supplementing the white cane. Provides 360° dual-elevation hazard perception directly mapped to torso azimuth."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_DARK

    p = rtf.add_paragraph()
    p.text = "\nEXPECTED MOBILITY BENEFIT"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = rtf.add_paragraph()
    p.text = "Targets a 35% speed increase (to 0.70 m/s) and zero head-level collisions. Device is designed as a supplementary aid—not replacing the cane."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_DARK

    stages = [
        ("1. SENSE", "ToF LiDARs measure distances\nacross 0° & +25° dual planes"),
        ("2. INTERPRET", "ESP32-S3 Core 0 DSP filters\nmeasurements & eliminates dropouts"),
        ("3. LOCALIZE", "Firmware identifies nearest hazard\n& associates radial waist node"),
        ("4. ALERT", "Core 1 modulates LRA cadence\n(1.2s to 12Hz) & directional intensity")
    ]
    for idx, (s_title, s_desc) in enumerate(stages):
        sx = Inches(0.4 + idx * 3.15)
        sc = create_card(slide4, sx, Inches(5.8), Inches(3.05), Inches(0.95), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
        stf = sc.text_frame
        stf.margin_left = stf.margin_right = stf.margin_top = Inches(0.1)
        stf.word_wrap = True
        p = stf.paragraphs[0]
        p.text = s_title
        p.font.name = FONT_BODY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(56, 189, 248)
        p = stf.add_paragraph()
        p.text = s_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 5: SYSTEM ARCHITECTURE
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "SYSTEM ARCHITECTURE", "HARDWARE & FIRMWARE PIPELINE")

    sub_box = slide5.shapes.add_textbox(Inches(0.4), Inches(0.75), Inches(12.5), Inches(0.3))
    stf = sub_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_top = 0
    sp = stf.paragraphs[0]
    sp.text = "DISTRIBUTED OPTICAL SENSING, DETERMINISTIC PROCESSING AND SPATIAL HAPTIC FEEDBACK"
    sp.font.name = FONT_BODY
    sp.font.size = Pt(9.5)
    sp.font.bold = True
    sp.font.color.rgb = BLUE_ACCENT

    if os.path.exists("assets/system_architecture_diagram.png"):
        slide5.shapes.add_picture("assets/system_architecture_diagram.png", Inches(0.4), Inches(1.1), width=Inches(9.2), height=Inches(5.7))

    arch_panel = create_card(slide5, Inches(9.8), Inches(1.1), Inches(3.1), Inches(5.7), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    atf = arch_panel.text_frame
    atf.margin_left = atf.margin_right = atf.margin_top = Inches(0.12)
    atf.word_wrap = True

    p = atf.paragraphs[0]
    p.text = "ARCHITECTURE SPECIFICATIONS"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.8)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = atf.add_paragraph()
    p.text = "\n1. Optical Perception Subsystem\n• 4x 0° waist sensors + 4x +25° head sensors\n• ST VL53L1X (940nm Class-1 eye-safe)\n• TCA9548A multiplexer with interleaved 4ms firing (15 Hz aggregate rate)"
    p.font.name = FONT_BODY
    p.font.size = Pt(7.2)
    p.font.color.rgb = TEXT_DARK

    p = atf.add_paragraph()
    p.text = "\n2. Central Processing (ESP32-S3)\n• Core 0: Non-blocking Fast-Mode I²C sensor polling & median filtering\n• Core 1: Distance-to-cadence mapping, spatial haptic interpolation & 3s alert decay\n• ULP Coprocessor: BMI160 step detection, 15µA light sleep, 5ms fast wake\n• Latency: Target <15 ms sensor-to-actuation"
    p.font.name = FONT_BODY
    p.font.size = Pt(7.2)
    p.font.color.rgb = TEXT_DARK

    p = atf.add_paragraph()
    p.text = "\n3. Resonant Haptic Actuation\n• 8x TI DRV2605L motor driver ICs\n• 8x Precision 205Hz LRAs (up to 2.2G peak)\n• 150mA hardware current limit"
    p.font.name = FONT_BODY
    p.font.size = Pt(7.2)
    p.font.color.rgb = TEXT_DARK

    p = atf.add_paragraph()
    p.text = "\n4. Power Management & Safety\n• 3.7V 2500mAh Li-ion (IEC 62133 reference)\n• Hardware BMS: 4.25V OVP, 2.8V UVP, 2A OCP\n• Skin contact target <34.5°C"
    p.font.name = FONT_BODY
    p.font.size = Pt(7.2)
    p.font.color.rgb = TEXT_DARK

    p = atf.add_paragraph()
    p.text = "\n5. Accessible Companion Platform\n• Offline BLE link; zero cloud dependency\n• TalkBack / VoiceOver accessible UI\n• O&M trainer 360° radar mirror"
    p.font.name = FONT_BODY
    p.font.size = Pt(7.2)
    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 6: TARGET SPECIFICATIONS & HARDWARE COMPONENTS
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "TARGET SPECIFICATIONS & HARDWARE COMPONENTS", "ENGINEERING TARGETS & SELECTION")

    t_a_shape = slide6.shapes.add_table(16, 3, Inches(0.4), Inches(0.85), Inches(7.2), Inches(6.1))
    table_a = t_a_shape.table
    table_a.columns[0].width = Inches(1.8)
    table_a.columns[1].width = Inches(2.2)
    table_a.columns[2].width = Inches(3.2)

    table_a.cell(0, 0).text = "PARAMETER"
    table_a.cell(0, 1).text = "TARGET SPECIFICATION"
    table_a.cell(0, 2).text = "IMPLEMENTATION APPROACH"

    specs_data = [
        ("Core Controller", "ESP32-S3 Dual-Core MCU, 240 MHz", "Real-time dual-core DSP; zero neural network latency"),
        ("Optical Sensing", "8x ST VL53L1X Time-of-Flight Sensors", "Class-1 940 nm infrared micro-LiDARs (27° FoV)"),
        ("Horizontal Sensing Plane", "4 sensors at 0° elevation", "Front, Rear, Left, Right waist-level obstacle detection"),
        ("Upper Sensing Plane", "4 sensors at +25° elevation", "Projects protective cone from chest level to 2.1 m height"),
        ("Angular Distribution", "8 positions at 45° intervals", "Complete 360° radial perimeter around user's waist"),
        ("Aggregate Refresh Rate", "15 Hz target aggregate rate", "Fast-Mode I²C polling across multiplexer channels"),
        ("Sensor Scheduling", "Interleaved sequence: 1→5→2→6→3→7→4→8", "4 ms phase offsets prevent optical crosstalk dropouts"),
        ("Processing Latency", "Under 15 ms target latency", "Hard real-time deterministic DSP on Core 0 & Core 1"),
        ("Haptic Actuation", "8x Linear Resonant Actuators (LRAs)", "Driven by 8x TI DRV2605L motor controllers via I²C"),
        ("Haptic Resonance", "205 Hz ± 3 Hz auto-tracking", "Locked to mechanical resonance; up to 2.2 G acceleration"),
        ("Proximity Response", "Distance-to-cadence repetition", "1.2 s pulse (2.5 m) accelerating to 12 Hz burst (0.4 m)"),
        ("Optical Field of View", "27° diagonal cone per sensor", "Specified in source documentation for VL53L1X optics"),
        ("Sensor Wavelength", "940 nm infrared (Class-1)", "IEC 60825-1:2014 eye-safe compliance reference"),
        ("Battery Subsystem", "3.7 V, 2500 mAh Li-ion rechargeable", "Hardware BMS (4.25V OVP, 2.8V UVP, 2A OCP, IEC 62133 ref)"),
        ("Prototype BOM Cost", "₹9,500 estimated prototype BOM", "Source BOM; ₹3,650 volume projection (1,000+ units)")
    ]

    for r_idx, (p1, p2, p3) in enumerate(specs_data, 1):
        table_a.cell(r_idx, 0).text = p1
        table_a.cell(r_idx, 1).text = p2
        table_a.cell(r_idx, 2).text = p3

    style_table(table_a, header_bg=NAVY_PRIMARY)

    t_b_shape = slide6.shapes.add_table(10, 3, Inches(7.8), Inches(0.85), Inches(5.1), Inches(6.1))
    table_b = t_b_shape.table
    table_b.columns[0].width = Inches(0.6)
    table_b.columns[1].width = Inches(2.0)
    table_b.columns[2].width = Inches(2.5)

    table_b.cell(0, 0).text = "S. NO."
    table_b.cell(0, 1).text = "HARDWARE COMPONENT"
    table_b.cell(0, 2).text = "PRIMARY USE"

    hw_data = [
        ("1", "ESP32-S3-WROOM-1 MCU", "Dual-core processing, sensor scheduling, haptic DSP, and BLE telemetry"),
        ("2", "8x ST VL53L1X ToF LiDARs", "360° environmental distance measurement across 0° and +25° planes"),
        ("3", "TCA9548A I²C Multiplexer", "Eliminates I²C address conflicts and manages interleaved sensor polling"),
        ("4", "8x TI DRV2605L Haptic ICs", "Provides auto-resonance tracking, 150 mA current limit, and crisp transients"),
        ("5", "8x 205 Hz LRA Coin Motors", "Delivers localized vibrotactile pressure gradients to abdominal wall"),
        ("6", "BMI160 6-DOF IMU", "Detects user movement; triggers 15 µA light sleep and 5 ms motion wake-up"),
        ("7", "3.7V 2500mAh Li-ion + BMS", "Portable rechargeable power with hardware over/under voltage protection"),
        ("8", "USB-C & BLE Interfaces", "Safe battery charging, firmware flashing, and accessible companion link"),
        ("9", "Neoprene Belt & TPU Pods", "Ergonomic wearable integration, IP54 ingress protection, and cable harness")
    ]

    for r_idx, (h1, h2, h3) in enumerate(hw_data, 1):
        table_b.cell(r_idx, 0).text = h1
        table_b.cell(r_idx, 1).text = h2
        table_b.cell(r_idx, 2).text = h3

    style_table(table_b, header_bg=BLUE_ACCENT)

    # ==========================================
    # SLIDE 7: CAD DESIGN & MECHANICAL SPECS
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "CAD DESIGN & MECHANICAL SPECS", "MECHANICAL PACKAGING & ERGONOMICS")

    if os.path.exists("assets/spatial_sense_worn.jpg"):
        slide7.shapes.add_picture("assets/spatial_sense_worn.jpg", Inches(0.4), Inches(0.85), width=Inches(6.1), height=Inches(3.3))
    
    cap1 = create_card(slide7, Inches(0.4), Inches(3.8), Inches(6.1), Inches(0.35), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
    c1_tf = cap1.text_frame
    c1_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = c1_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Assembled CAD Concept: Ergonomic Waistband with 8 Modular Dual-Elevation Pods"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.0)
    p.font.color.rgb = WHITE

    if os.path.exists("assets/spatial_sense_exploded.jpg"):
        slide7.shapes.add_picture("assets/spatial_sense_exploded.jpg", Inches(6.8), Inches(0.85), width=Inches(6.1), height=Inches(3.3))

    cap2 = create_card(slide7, Inches(6.8), Inches(3.8), Inches(6.1), Inches(0.35), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
    c2_tf = cap2.text_frame
    c2_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = c2_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Exploded Assembly: Modular TPU Pods, ToF Sensors, LRAs, FPC Harness & Enclosure"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.0)
    p.font.color.rgb = WHITE

    p_a = create_card(slide7, Inches(0.4), Inches(4.3), Inches(6.1), Inches(2.7), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    atf = p_a.text_frame
    atf.margin_left = atf.margin_right = atf.margin_top = Inches(0.12)
    atf.word_wrap = True
    p = atf.paragraphs[0]
    p.text = "ENCLOSURE & WEARABLE INTEGRATION"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p = atf.add_paragraph()
    p.text = "• Ergonomic Belt: Breathable, contoured neoprene belt with micro-velcro adjustment (waist 26–44 in), keeping total wearable weight under 320 grams.\n• Modular Sensor Pods: Eight injection-molded TPU pods with internal silicone dampening gaskets to isolate LRA vibration from optical ToF sensor cavities.\n• Flexible Harness: 2-layer polyimide flexible printed circuit (FPC) routing power and I²C buses with zero exposed wires, eliminating clothing snag risks.\n• Electronics & Battery Housing: Contoured rear enclosure securing the ESP32-S3 board, BMS, and 2500 mAh cell with magnetic alignment.\n• Accessible Interface: Tactile raised hardware self-test button and recessed USB-C charge port operable entirely without sight.\n• Serviceability: Snap-fit pod shells permit individual sensor or actuator replacement in the field without replacing the entire belt assembly."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.5)
    p.font.color.rgb = TEXT_DARK

    p_b = create_card(slide7, Inches(6.8), Inches(4.3), Inches(6.1), Inches(2.7), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    btf = p_b.text_frame
    btf.margin_left = btf.margin_right = btf.margin_top = Inches(0.12)
    btf.word_wrap = True
    p = btf.paragraphs[0]
    p.text = "MECHANICAL DESIGN & VALIDATION PROTOCOLS"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    p = btf.add_paragraph()
    p.text = "• Ergonomic Fit Assessment: User trials measuring torso contact pressure across walking, sitting, and bending to prevent localized chafing.\n• Optical Geometry Alignment: Precision jig verification ensuring 0° sensors remain within ±1.5° horizontal and +25° sensors maintain vertical cone.\n• Dynamic Gait Pod Stability: Accelerometer-based shake tests verifying pod retention and consistent skin-to-LRA acoustic coupling during walking.\n• Cable Strain Relief: 5,000-cycle cyclic bending test on FPC pod junctions at 90° radius to verify copper fatigue resilience.\n• Environmental Resistance: Ingress testing planned against IP54 dust and omnidirectional water splash standards.\n• Thermal Profiling: Thermocouple measurement targeting skin-facing contact surface <34.5°C under continuous 60-min resonance actuation.\n• Accessibility Verification: Blindfolded donning, adjustment, and buckle fastening evaluation conducted with assistive technology specialists."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.5)
    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 8: ELECTRONICS SCHEMATICS & BOM
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "ELECTRONICS SCHEMATICS & BOM", "ELECTRICAL SCHEMATICS & BILL OF MATERIALS")

    schems = [
        ("assets/schematic_sheet1.png", Inches(0.4), Inches(0.85), Inches(3.3), Inches(2.3)),
        ("assets/schematic_sheet2.png", Inches(3.9), Inches(0.85), Inches(3.3), Inches(2.3)),
        ("assets/schematic_sheet3.png", Inches(0.4), Inches(3.25), Inches(3.3), Inches(2.3)),
        ("assets/schematic_sheet4.png", Inches(3.9), Inches(3.25), Inches(3.3), Inches(2.3))
    ]
    for s_path, sx, sy, sw, sh in schems:
        if os.path.exists(s_path):
            slide8.shapes.add_picture(s_path, sx, sy, width=sw, height=sh)

    s_note = create_card(slide8, Inches(0.4), Inches(5.65), Inches(6.8), Inches(1.35), bg_color=RGBColor(241, 245, 249), border_color=BLUE_CARD_BORDER)
    sntf = s_note.text_frame
    sntf.margin_left = sntf.margin_right = sntf.margin_top = Inches(0.1)
    sntf.word_wrap = True
    p = sntf.paragraphs[0]
    p.text = "ELECTRICAL ARCHITECTURE NOTES"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.0)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p = sntf.add_paragraph()
    p.text = "• Sheet 1: Fast-Mode I²C Bus 0 (400 kHz) interfaces TCA9548A mux to 8x VL53L1X ToF sensors with individual XSHUT control lines.\n• Sheet 2: Independent I²C Bus 1 interfaces 8x TI DRV2605L drivers; auto-tracks 205 Hz resonance with 150 mA hardware current limit.\n• Sheet 3: 3.7V 2500mAh Li-ion battery with hardware BMS (4.25V OVP, 2.8V UVP, 2A OCP) and low-noise 3.3V buck-boost regulation.\n• Sheet 4: BMI160 6-DOF IMU triggers ULP 15µA light sleep / 5ms wake-up; BLE 5.0 provides offline-first TalkBack/VoiceOver telemetry."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.0)
    p.font.color.rgb = TEXT_DARK

    bom_shape = slide8.shapes.add_table(7, 4, Inches(7.5), Inches(0.85), Inches(5.4), Inches(3.8))
    table_bom = bom_shape.table
    table_bom.columns[0].width = Inches(1.8)
    table_bom.columns[1].width = Inches(1.8)
    table_bom.columns[2].width = Inches(0.8)
    table_bom.columns[3].width = Inches(1.0)

    table_bom.cell(0, 0).text = "SUBSYSTEM"
    table_bom.cell(0, 1).text = "PART DESIGNATION"
    table_bom.cell(0, 2).text = "QTY"
    table_bom.cell(0, 3).text = "EST. COST"

    bom_items = [
        ("Core Processing & Power", "ESP32-S3-WROOM-1 + BMS + 2500mAh Li-ion", "1 Set", "₹1,450"),
        ("Optical Sensing Array", "ST VL53L1X ToF Sensors + Optics", "8 Units", "₹3,200"),
        ("Multiplexed Haptic Drivers", "TI DRV2605L + TCA9548A Mux PCB", "8 Drivers + Mux", "₹2,100"),
        ("Resonant Actuators", "Precision 205Hz LRA Coin Actuators", "8 Units", "₹1,250"),
        ("Harness, Enclosure & Belt", "Modular TPU Pods, Neoprene Belt & FPC", "1 Assembly", "₹1,500"),
        ("TOTAL PROTOTYPE BOM", "Source-Stated Functional Prototype", "—", "₹9,500")
    ]

    for r_idx, (b1, b2, b3, b4) in enumerate(bom_items, 1):
        table_bom.cell(r_idx, 0).text = b1
        table_bom.cell(r_idx, 1).text = b2
        table_bom.cell(r_idx, 2).text = b3
        table_bom.cell(r_idx, 3).text = b4

    style_table(table_bom, header_bg=NAVY_PRIMARY)
    for c_idx in range(4):
        cell = table_bom.cell(6, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(239, 246, 255)
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = NAVY_PRIMARY

    econ_card = create_card(slide8, Inches(7.5), Inches(4.8), Inches(5.4), Inches(2.2), bg_color=WHITE, border_color=GREEN_CARD_BORDER)
    etf = econ_card.text_frame
    etf.margin_left = etf.margin_right = etf.margin_top = Inches(0.12)
    etf.word_wrap = True
    p = etf.paragraphs[0]
    p.text = "UNIT ECONOMICS & SCALABILITY ANALYSIS"
    p.font.name = FONT_BODY
    p.font.size = Pt(8.8)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    p = etf.add_paragraph()
    p.text = "• Functional Prototype BOM: ₹9,500 (component subtotal for single bench/wearable prototype).\n• Commercial Volume Manufacturing: Projected to reduce to ₹3,650 per unit at 1,000+ units through automated SMT assembly, reel-based component sourcing, and injection-molded TPU tooling.\n• Target Retail Price: Benchmarked at ₹4,500 retail, offering a transformative alternative to imported haptic vests (₹1,20,000+) and handheld smart canes (₹4,000–₹6,000).\n• Subsidized Distribution: Designed for integration into government welfare schemes (ADIP / ALIMCO) and partnership distribution through EnAble India and Blind Persons Associations."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.5)
    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 9: MAKER JOURNEY
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "MAKER JOURNEY", "STAGE 1 TO STAGE 2 ITERATIONS")

    sub_box = slide9.shapes.add_textbox(Inches(0.4), Inches(0.75), Inches(12.5), Inches(0.3))
    stf = sub_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_top = 0
    sp = stf.paragraphs[0]
    sp.text = "FROM SYSTEM CONCEPT TO WEARABLE PROTOTYPE • HARDWARE & FIRMWARE ITERATIONS"
    sp.font.name = FONT_BODY
    sp.font.size = Pt(9.5)
    sp.font.bold = True
    sp.font.color.rgb = BLUE_ACCENT

    if os.path.exists("assets/workbench_lab.jpg"):
        slide9.shapes.add_picture("assets/workbench_lab.jpg", Inches(0.4), Inches(1.1), width=Inches(3.8), height=Inches(2.5))
    cap_f1 = create_card(slide9, Inches(0.4), Inches(3.45), Inches(3.8), Inches(0.45), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
    c1_tf = cap_f1.text_frame
    c1_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = c1_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Fig 1: Benchtop prototype rig: ESP32-S3 wired to ToF LiDARs, LRAs, and oscilloscope showing PWM waveforms."
    p.font.name = FONT_BODY
    p.font.size = Pt(6.8)
    p.font.color.rgb = WHITE

    if os.path.exists("assets/cad_ui_feature_tree.jpg"):
        slide9.shapes.add_picture("assets/cad_ui_feature_tree.jpg", Inches(4.4), Inches(1.1), width=Inches(3.8), height=Inches(2.5))
    cap_f2 = create_card(slide9, Inches(4.4), Inches(3.45), Inches(3.8), Inches(0.45), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
    c2_tf = cap_f2.text_frame
    c2_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = c2_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Fig 2: 3D CAD modeling UI showing active parametric feature tree for the modular TPU pod and belt mount."
    p.font.name = FONT_BODY
    p.font.size = Pt(6.8)
    p.font.color.rgb = WHITE

    if os.path.exists("assets/eda_kicad_layout.jpg"):
        slide9.shapes.add_picture("assets/eda_kicad_layout.jpg", Inches(0.4), Inches(4.05), width=Inches(7.8), height=Inches(2.45))
    cap_f3 = create_card(slide9, Inches(0.4), Inches(6.55), Inches(7.8), Inches(0.45), bg_color=NAVY_PRIMARY, border_color=NAVY_PRIMARY)
    c3_tf = cap_f3.text_frame
    c3_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = c3_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Fig 3: EDA layout software (KiCad) showing 2-layer polyimide flexible PCB harness routing with 0-error DRC rule-check logs."
    p.font.name = FONT_BODY
    p.font.size = Pt(7.0)
    p.font.color.rgb = WHITE

    iter_panel = create_card(slide9, Inches(8.4), Inches(1.1), Inches(4.5), Inches(5.9), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    itf = iter_panel.text_frame
    itf.margin_left = itf.margin_right = itf.margin_top = Inches(0.15)
    itf.word_wrap = True

    p = itf.paragraphs[0]
    p.text = "STAGE 1 TO STAGE 2 ENGINEERING ITERATIONS"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    iterations = [
        ("1. Optical Sensing Geometry Evolution",
         "Shifted from single-plane waist detection to dual-elevation geometry (4x 0° waist + 4x +25° upward sensors). This creates a 2.1m protective cone to detect overhead tree limbs and shop awnings without increasing sensor count."),
        ("2. Deterministic Crosstalk Mitigation",
         "Developed an interleaved firing sequence (1→5→2→6→3→7→4→8) with 4 ms phase offsets managed via TCA9548A multiplexer. Eliminates optical interference while sustaining a 15 Hz aggregate refresh rate."),
        ("3. Resonant Haptic Actuator Selection",
         "Replaced generic ERM vibration motors with precision 205 Hz Linear Resonant Actuators (LRAs) driven by TI DRV2605L controllers. Achieves 5 ms rapid rise time, 2.2 G peak acceleration, and 60% lower power."),
        ("4. Flexible PCB & Pod Packaging",
         "Transitioned from bulky discrete wiring to a custom 2-layer polyimide flex-PCB (FPC) harness housed within injection-molded TPU pods and breathable neoprene belt, reducing total wearable weight to <320g."),
        ("5. Power & Thermal Management",
         "Integrated BMI160 6-DOF IMU to trigger 15 µA light sleep when stationary with 5 ms fast wake upon step detection. Verified skin-facing pod temperature target <34.5°C under continuous resonance."),
        ("6. Inclusive Deaf-Blind Accessibility",
         "Implemented a zero-sight physical belt self-test (double-tap triggers clockwise verification sweep) and an offline BLE companion application compliant with TalkBack/VoiceOver and O&M trainer radar mirroring.")
    ]

    for i_title, i_text in iterations:
        p = itf.add_paragraph()
        p.text = f"\n{i_title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(7.8)
        p.font.bold = True
        p.font.color.rgb = BLUE_ACCENT

        p = itf.add_paragraph()
        p.text = i_text
        p.font.name = FONT_BODY
        p.font.size = Pt(7.0)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 10: SAFETY & HAZARD MITIGATION
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "SAFETY & HAZARD MITIGATION", "SAFETY & COMPLIANCE")

    safe_panel = create_card(slide10, Inches(0.4), Inches(0.85), Inches(6.1), Inches(6.1), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    stf = safe_panel.text_frame
    stf.margin_left = stf.margin_right = stf.margin_top = Inches(0.18)
    stf.word_wrap = True

    p = stf.paragraphs[0]
    p.text = "⬟ SAFETY & HAZARD MITIGATION MEASURES"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    safety_items = [
        ("1. Electrical & Battery Safety",
         "• Powered by an IEC 62133-certified 3.7V 2500 mAh Li-ion cell.\n• Hardware BMS provides multi-tier protection: 4.25V over-voltage (OVP), 2.80V under-voltage (UVP), and 2.0A over-current (OCP).\n• DRV2605L driver current is hardware-limited to 150 mA, preventing actuator drive circuit overload.", NAVY_PRIMARY),
        ("2. Thermal Safety & Contact Temperature",
         "• Under continuous resonant actuation for 60 minutes, the skin-facing pod surface temperature is targeted to remain below 34.5°C.\n• Well within safe human skin contact threshold (<41°C per medical wearable guidelines), preventing localized burns or discomfort.", AMBER_ACCENT),
        ("3. Optical / Ocular Radiation Safety",
         "• All 8 ST VL53L1X optical emitters comply with Class 1 Laser Safety standards (IEC 60825-1:2014) at 940 nm wavelength.\n• Poses zero ocular hazard to user, bystanders, or guide animals under all operating conditions.", GREEN_ACCENT),
        ("4. Sensory Fatigue Prevention",
         "• Continuous tactile stimulation induces mechanoreceptor desensitization within 45 seconds.\n• Firmware enforces a 3-second alert decay window: if user pauses before an obstacle, pulse intensity gracefully decays to maintain sensitivity.", BLUE_ACCENT),
        ("5. Wearability & Mechanical Safety",
         "• Rounded pod geometry prevents clothing snagging; internal silicone dampeners prevent vibration bleed to frame.\n• Flexible strain-relieved FPC prevents skin pinching; quick-release magnetic/velcro buckle ensures rapid removal.", NAVY_PRIMARY)
    ]

    for s_title, s_desc, s_col in safety_items:
        p = stf.add_paragraph()
        p.text = f"\n{s_title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = s_col
        p = stf.add_paragraph()
        p.text = s_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    std_panel = create_card(slide10, Inches(6.8), Inches(0.85), Inches(6.1), Inches(6.1), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    stf2 = std_panel.text_frame
    stf2.margin_left = stf2.margin_right = stf2.margin_top = Inches(0.18)
    stf2.word_wrap = True

    p = stf2.paragraphs[0]
    p.text = "✦ STANDARDS & FIELD COMPLIANCE MATRIX"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    standards_items = [
        ("IEC 62133-2:2017 Battery Safety Reference",
         "Secondary lithium cells and batteries for portable applications. Enforces thermal shock, short-circuit, overcharge, and mechanical drop compliance criteria."),
        ("IEC 60825-1:2014 Laser Eye-Safety Classification",
         "Safety of laser products: 940 nm infrared VCSEL emitters documented as Class 1 inherently eye-safe under all magnification and exposure conditions."),
        ("Thermal Contact & Skin Protection Verification",
         "Planned thermal camera and thermocouple profiling under sustained 100% actuation duty cycle to document skin-facing temperature below 34.5°C."),
        ("IP54 Enclosure Ingress Protection",
         "Mechanical enclosure designed to prevent harmful dust ingress and resist multi-directional splashing water during rain or outdoor use."),
        ("Electromagnetic Compatibility (EMC / RED)",
         "Verification of 2.4 GHz BLE emissions and low-noise DC-DC regulator switching to prevent interference with medical pacemakers or hearing aids."),
        ("Accessibility Standards & Assistive Compliance",
         "Companion mobile interface evaluated against Google Accessibility Scanner and W3C WCAG 2.1 AA for TalkBack and VoiceOver compliance."),
        ("Structured Controlled Mobility Trials",
         "Multi-stage obstacle courses conducted with blindfolded volunteers and institutional partners (EnAble India) under certified O&M trainer supervision."),
        ("Documented FMEA & Fail-Safe Protocols",
         "Formal failure modes and effects analysis with automatic system fallbacks, watchdog supervisor, and accessible non-visual fault notifications.")
    ]

    for std_title, std_desc in standards_items:
        p = stf2.add_paragraph()
        p.text = f"\n• {std_title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = stf2.add_paragraph()
        p.text = std_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 11: ROADMAP & FMEA
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "ROADMAP & FAILURE MODE AND EFFECT ANALYSIS", "RELIABILITY & MILESTONES")

    fmea_shape = slide11.shapes.add_table(7, 3, Inches(0.4), Inches(0.85), Inches(12.5), Inches(3.2))
    table_fmea = fmea_shape.table
    table_fmea.columns[0].width = Inches(3.0)
    table_fmea.columns[1].width = Inches(4.3)
    table_fmea.columns[2].width = Inches(5.2)

    table_fmea.cell(0, 0).text = "FAILURE MODE"
    table_fmea.cell(0, 1).text = "POTENTIAL HAZARD / EFFECT"
    table_fmea.cell(0, 2).text = "PROPOSED FAIL-SAFE / MITIGATION"

    fmea_data = [
        ("Optical Crosstalk or Invalid Distance", 
         "False obstacle alerts or missed hazards leading to user hesitation or impact.", 
         "Interleaved firing sequence (4 ms phase offsets), temporal median filtering, timeout error traps, and optical bench validation."),
        ("Sensor Obstruction or Dirt Accumulation", 
         "Reduced detection coverage in specific azimuth sector without user awareness.", 
         "Per-channel SNR / ambient light monitoring, sector dropout detection, and accessible non-visual fault indicator pulse."),
        ("Incorrect Directional Node Mapping", 
         "Misleading tactile direction communicated, prompting user to veer toward hazard.", 
         "Factory azimuth calibration table, spatial haptic interpolation verification, and real-time companion app radar mirror check."),
        ("Haptic Actuator / Driver Failure", 
         "Obstacle detected optically but warning not perceived physically by user.", 
         "Startup self-test (clockwise verification sweep), DRV2605L fault register monitoring, and adjacent-node fallback alerting."),
        ("Battery Depletion / Sudden Cutoff", 
         "Abrupt loss of navigation protection while user is walking outdoors.", 
         "Early 15% low-battery distinctive 3-pulse cadence warning, power-saving light sleep, and controlled safe shutdown at 2.8V."),
        ("Controller Firmware Stall / Bus Hang", 
         "Stale sensor data or continuous locked vibration causing severe confusion.", 
         "Hardware Watchdog Timer (WDT) with auto-reset, I²C bus timeout clear routines, and normally-quiescent actuator fail-safe.")
    ]

    for r_idx, (f1, f2, f3) in enumerate(fmea_data, 1):
        table_fmea.cell(r_idx, 0).text = f1
        table_fmea.cell(r_idx, 1).text = f2
        table_fmea.cell(r_idx, 2).text = f3

    style_table(table_fmea, header_bg=NAVY_PRIMARY)

    phases = [
        ("PHASE 1 (MONTHS 1–2)", "Benchtop Validation & Crosstalk Quantification", 
         "• Breadboard integration of 8x ToF sensors + TCA9548A mux.\n• Timing calibration & interleaved phase offset verification.\n• Crosstalk quantification across varied optical surfaces.\n• Deliverable: Bench rig achieving 15 Hz aggregate rate with 0 false dropouts.", BLUE_ACCENT),
        ("PHASE 2 (MONTHS 3–4)", "Custom Flex-PCB & Ergonomic Pod Enclosure", 
         "• Fabricate 2-layer polyimide flexible PCB (FPC) harness.\n• Tooling and casting of modular IP54 silicone-lined TPU pods.\n• Tune DRV2605L resonance frequencies to 205 Hz ± 3 Hz.\n• Deliverable: Populated wearable prototype; thermal profile <34.5°C.", NAVY_PRIMARY),
        ("PHASE 3 (MONTH 5)", "Firmware Refinement & BLE Companion App", 
         "• Implement distance-to-cadence mapping (1.2s to 12Hz).\n• Test virtual-point spatial haptic interpolation between nodes.\n• Integrate TalkBack / VoiceOver offline BLE interface.\n• Deliverable: Feature-complete belt; app passing Accessibility Scanner.", GREEN_ACCENT),
        ("PHASE 4 (MONTH 6)", "Pilot User Trials & EnAble India Validation", 
         "• Structured indoor and outdoor obstacle navigation trials.\n• Blindfolded volunteers and institutional beta testers.\n• Quantitative metrics: head-strike avoidance and transit speed.\n• Deliverable: Formal pilot trial evaluation report with EnAble India.", AMBER_ACCENT)
    ]

    for idx, (p_num, p_title, p_desc, p_col) in enumerate(phases):
        px = Inches(0.4 + idx * 3.15)
        pc = create_card(slide11, px, Inches(4.25), Inches(3.05), Inches(2.8), bg_color=WHITE, border_color=p_col, line_width=1.5)
        ptf = pc.text_frame
        ptf.margin_left = ptf.margin_right = ptf.margin_top = Inches(0.12)
        ptf.word_wrap = True
        p = ptf.paragraphs[0]
        p.text = p_num
        p.font.name = FONT_BODY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = p_col
        p = ptf.add_paragraph()
        p.text = p_title
        p.font.name = FONT_TITLE
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = ptf.add_paragraph()
        p.text = p_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 12: TECHNICAL CHALLENGES & RISKS
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "TECHNICAL CHALLENGES & RISKS", "RISK MATRIX & FAIL-SAFES")

    left_bot = create_card(slide12, Inches(0.4), Inches(0.85), Inches(4.8), Inches(6.1), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    ltf = left_bot.text_frame
    ltf.margin_left = ltf.margin_right = ltf.margin_top = Inches(0.18)
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "ENGINEERING BOTTLENECKS & CHALLENGES"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    bottlenecks = [
        ("• Multi-Sensor Optical Crosstalk", "Mitigating overlapping 940 nm laser reflections across eight closely spaced ToF micro-LiDARs operating in simultaneous proximity."),
        ("• Varied Surface Reflectance", "Maintaining accurate photon ranging across dark fabrics, transparent glass partitions, wet pavements, and diffuse vegetation."),
        ("• Real-Time Processing Latency", "Executing non-blocking I²C polling, multiplexer switching, and haptic PWM modulation under strict <15 ms latency without OS jitter."),
        ("• Spatial Haptic Localization", "Ensuring user distinctly perceives eight discrete radial torso vectors without tactile blending or confusion between adjacent nodes."),
        ("• Mechanoreceptor Sensory Adaptation", "Preventing human cutaneous tactile fatigue under sustained proximity warnings without muting safety-critical alerts."),
        ("• Thermal & Power Constraints", "Delivering up to 2.2 G tactile pulses while keeping skin contact surface <34.5°C and sustaining all-day wearable battery life."),
        ("• Non-Visual Deaf-Blind Accessibility", "Designing 100% accessible physical controls and offline status indication operable with zero auditory or visual cues.")
    ]

    for b_title, b_desc in bottlenecks:
        p = ltf.add_paragraph()
        p.text = f"\n{b_title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = BLUE_ACCENT
        p = ltf.add_paragraph()
        p.text = b_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    r_top_shape = slide12.shapes.add_table(7, 2, Inches(5.5), Inches(0.85), Inches(7.4), Inches(2.9))
    t_rt = r_top_shape.table
    t_rt.columns[0].width = Inches(2.5)
    t_rt.columns[1].width = Inches(4.9)

    t_rt.cell(0, 0).text = "TECHNICAL RISK"
    t_rt.cell(0, 1).text = "PROPOSED MITIGATION"

    rt_data = [
        ("Optical Crosstalk Between ToF Sensors", "Interleaved firing sequence (1→5→2→6→3→7→4→8) with 4 ms phase offsets."),
        ("Incorrect Overhead Hazard Localization", "Calibrate dual-plane geometry (0° vs +25°) and verify radial ray-trace angles."),
        ("Tactile Cadence Misinterpretation", "Distance-to-cadence tuning (1.2s to 12Hz) with distinct emergency burst at 0.4m."),
        ("I²C Communication Bus Stall / Latch", "Implement bus timeout recovery routines and TCA9548A hardware reset pin."),
        ("Battery Voltage Sag Under Peak Load", "Low-ESR decoupling capacitors, independent 3.3V rails, and BMS current limiting."),
        ("Actuator Pod Overheating on Skin", "Active duty-cycle management, thermal isolation silicone, and hardware 150mA limit.")
    ]

    for r_idx, (r1, r2) in enumerate(rt_data, 1):
        t_rt.cell(r_idx, 0).text = r1
        t_rt.cell(r_idx, 1).text = r2

    style_table(t_rt, header_bg=NAVY_PRIMARY)

    r_bot_shape = slide12.shapes.add_table(7, 2, Inches(5.5), Inches(3.95), Inches(7.4), Inches(2.95))
    t_rb = r_bot_shape.table
    t_rb.columns[0].width = Inches(2.5)
    t_rb.columns[1].width = Inches(4.9)

    t_rb.cell(0, 0).text = "DESIGN CHALLENGE"
    t_rb.cell(0, 1).text = "PROPOSED FAIL-SAFE / ALTERNATIVE STRATEGY"

    rb_data = [
        ("Sensor Noise / False Distance Dropouts", "Temporal median filtering, outlier rejection, and multi-frame consistency checks."),
        ("LRA Actuator Hardware Disconnection", "Automatic power-on self-test sweep; fail-over warning via adjacent healthy nodes."),
        ("Ambiguous Tactile Proximity Gradient", "Strict non-linear cadence mapping (1.2s, 0.6s, 0.2s, 12Hz burst) verified in trials."),
        ("Improper Belt Fit or Sensor Rotation", "Contoured neoprene strap with non-slip silicone inner lining and IMU orientation check."),
        ("Sudden Power Loss During Walking", "Early triple-pulse low-battery warning and non-volatile parameter saving."),
        ("Field Repair & Maintenance Difficulty", "Modular snap-fit TPU pods permitting single-sensor or LRA replacement in field.")
    ]

    for r_idx, (r1, r2) in enumerate(rb_data, 1):
        t_rb.cell(r_idx, 0).text = r1
        t_rb.cell(r_idx, 1).text = r2

    style_table(t_rb, header_bg=BLUE_ACCENT)

    # ==========================================
    # SLIDE 13: REFERENCES & SOURCES
    # ==========================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "REFERENCES & SOURCES", "BIBLIOGRAPHY & TECHNICAL SOURCES")

    ref_left = create_card(slide13, Inches(0.4), Inches(0.85), Inches(6.1), Inches(6.1), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    rl_tf = ref_left.text_frame
    rl_tf.margin_left = rl_tf.margin_right = rl_tf.margin_top = Inches(0.18)
    rl_tf.word_wrap = True

    p = rl_tf.paragraphs[0]
    p.text = "COMPETITION PROPOSAL & ASSISTIVE CONTEXT"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    left_refs = [
        ("[1] Primary Competition Source:", 
         "Team TECHTONICS, 'SPATIAL-SENSE: A 360° Wearable Haptic-Sonar Waistband for Upper-Torso and Head-Level Collision Protection in Deaf-Blind Mobility', Vishwakarma Awards 2026–27 Stage 2 Proposal."),
        ("[2] Competition Framework & Mentorship:", 
         "Vishwakarma Awards 2026–27, Theme 2 — Assistive Solutions, organized in collaboration with EnAble India and Indian Institute of Technology Indore (IIT Indore)."),
        ("[3] Target Beneficiary Demographics:", 
         "National Institute for the Empowerment of Persons with Multiple Disabilities (NIEPMD) & Census of India: Demographic analysis indicating 450,000+ individuals with combined visual and hearing impairment in India."),
        ("[4] Orientation & Mobility (O&M) Standards:", 
         "EnAble India & Vision Aid Assistive Mobility Guidelines: White cane sweep mechanics, upper-torso collision biomechanics, and tactile interface standards for deaf-blind travelers."),
        ("[5] Tactile Sensory Adaptation Research:", 
         "Psychophysical studies on human cutaneous mechanoreceptor fatigue under sustained vibrotactile stimulation; foundation for the 3-second alert decay window.")
    ]

    for r_head, r_body in left_refs:
        p = rl_tf.add_paragraph()
        p.text = f"\n{r_head}"
        p.font.name = FONT_BODY
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = BLUE_ACCENT
        p = rl_tf.add_paragraph()
        p.text = r_body
        p.font.name = FONT_BODY
        p.font.size = Pt(7.2)
        p.font.color.rgb = TEXT_DARK

    if os.path.exists("assets/dual_plane_geometry.png"):
        slide13.shapes.add_picture("assets/dual_plane_geometry.png", Inches(0.6), Inches(5.1), width=Inches(5.7), height=Inches(1.6))

    ref_right = create_card(slide13, Inches(6.8), Inches(0.85), Inches(6.1), Inches(6.1), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
    rr_tf = ref_right.text_frame
    rr_tf.margin_left = rr_tf.margin_right = rr_tf.margin_top = Inches(0.18)
    rr_tf.word_wrap = True

    p = rr_tf.paragraphs[0]
    p.text = "HARDWARE DATASHEETS & SAFETY COMPLIANCE STANDARDS"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    right_refs = [
        ("[6] ESP32-S3 Dual-Core Controller Datasheet:", 
         "Espressif Systems, 'ESP32-S3 Series SoC Datasheet & Technical Reference Manual' — 240 MHz Xtensa dual-core 32-bit LX7, Fast-Mode I²C, ULP coprocessor and BLE 5.0 radio specifications."),
        ("[7] ST VL53L1X Time-of-Flight Ranging Sensor:", 
         "STMicroelectronics, 'VL53L1X Long Distance Ranging Time-of-Flight Sensor Datasheet' (DocID031281 Rev 4) — 940 nm VCSEL emitter, 27° FoV, Class 1 eye-safe optical ranging to 4.0 m."),
        ("[8] TI DRV2605L Haptic Driver Controller:", 
         "Texas Instruments, 'DRV2605L Low-Voltage Haptic Driver for LRA and ERM Actuators' (SLOS855) — Auto-resonance tracking, 150 mA current limit, and crisp 5 ms transient braking."),
        ("[9] TCA9548A 8-Channel I²C Multiplexer:", 
         "Texas Instruments, 'TCA9548A Low-Voltage 8-Channel I²C Switch with Reset' (SCPS144E) — Channel isolation, address conflict resolution, and hardware bus reset support."),
        ("[10] Bosch BMI160 6-DOF Inertial Measurement Unit:", 
         "Bosch Sensortec, 'BMI160 Small, Low Power 6-Axis IMU Datasheet' (BST-BMI160-DS000) — Low-power step detection and fast wake-up interrupt generation."),
        ("[11] IEC 62133-2:2017 Secondary Battery Safety Standard:", 
         "International Electrotechnical Commission, 'Secondary cells and batteries containing alkaline or other non-acid electrolytes — Portable lithium systems safety requirements.'"),
        ("[12] IEC 60825-1:2014 Laser Radiation Safety Standard:", 
         "International Electrotechnical Commission, 'Safety of laser products — Part 1: Equipment classification and requirements for Class 1 eye-safe optical emitters.'")
    ]

    for r_head, r_body in right_refs:
        p = rr_tf.add_paragraph()
        p.text = f"\n{r_head}"
        p.font.name = FONT_BODY
        p.font.size = Pt(7.8)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = rr_tf.add_paragraph()
        p.text = r_body
        p.font.name = FONT_BODY
        p.font.size = Pt(7.0)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 14: IMAGE SOURCES
    # ==========================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_header(slide14, "IMAGE SOURCES", "VISUAL DOCUMENTATION & CREDITS")

    credits_data = [
        ("PROJECT SOURCE & MOBILITY SCENARIO VISUALS",
         "• Figure 1: System Hardware Architecture & Dual-Plane Ray-Trace Geometry — Extracted from SPATIAL-SENSE Theme 2 Proposal (Page 2).\n• Figure 2: Resonant Haptic Distance-to-Cadence Timing Waveforms — Extracted from SPATIAL-SENSE Theme 2 Proposal (Page 3).\n• Figures 3 & 4: Before (Cane Alone Collision) & After (SPATIAL-SENSE Detection) Walking Scenarios — Extracted from SPATIAL-SENSE Theme 2 Proposal (Page 7).\n• Source Reference: Vishwakarma Awards 2026–27, Theme 2 Proposal Document.", BLUE_ACCENT),
        ("CAD MODELING, SCHEMATICS & EDA VISUALS",
         "• Assembled Wearable Concept & Exploded CAD Assembly — Parametric 3D CAD models generated for SPATIAL-SENSE Stage 2 submission by Team TECHTONICS.\n• Parametric CAD UI Interface Screenshot — Modeled representation of modular TPU sensor pod and snap-fit belt clip development.\n• Flexible PCB (FPC) Layout & DRC Check — Custom 2-layer polyimide flexible printed circuit routing designed in KiCad EDA.\n• Hardware Electrical Schematics — Sheets 1–4 functional block diagrams generated for Stage 2 submission.", NAVY_PRIMARY),
        ("INSTITUTIONAL BRANDING & CONCEPT VISUALIZATIONS",
         "• Vishwakarma Awards Official Logo — Maker Bhavan Foundation / Vishwakarma Awards 2026–27 Submission Kit.\n• Maker Bhavan Foundation Official Logo — Competition Template Asset.\n• Product Concept Photographs — AI-assisted conceptual visualization representing the proposed wearable form factor on user torso.\n• Workbench Hardware Rig — Laboratory benchtop development environment photograph representing prototype integration.", GREEN_ACCENT)
    ]

    for idx, (c_head, c_body, c_color) in enumerate(credits_data):
        cy = Inches(1.1 + idx * 1.85)
        cc = create_card(slide14, Inches(0.6), cy, Inches(12.133), Inches(1.65), bg_color=WHITE, border_color=BLUE_CARD_BORDER)
        ctf = cc.text_frame
        ctf.margin_left = ctf.margin_right = ctf.margin_top = Inches(0.18)
        ctf.word_wrap = True
        p = ctf.paragraphs[0]
        p.text = c_head
        p.font.name = FONT_BODY
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = c_color
        p = ctf.add_paragraph()
        p.text = c_body
        p.font.name = FONT_BODY
        p.font.size = Pt(8.0)
        p.font.color.rgb = TEXT_DARK

    out_path = "SPATIAL_SENSE_Theme2_Vishwakarma_2026_27.pptx"
    prs.save(out_path)
    print(f"Presentation successfully created: {out_path} ({len(prs.slides)} slides)")
    return out_path

if __name__ == "__main__":
    build_presentation()
